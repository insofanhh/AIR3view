"""Durable, project-backed production batches.

The browser never owns the queue. Each item is a normal AIR3view project and
uses the same `all` job as a manually imported video.
"""
import json
import io
import re
import shutil
import threading
import time
import uuid
import zipfile
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from . import store
from .models import Settings

MAX_ITEMS = 200
REFERENCE_EXTENSIONS = {'.wav', '.mp3', '.m4a', '.ogg', '.flac'}
MAX_REFERENCE_BYTES = 50 * 1024**2
MAX_ACTIVE = 2
MAX_AUTO_RETRIES = 5
AUTO_RETRY_DELAY = 3
MAX_TRANSIENT_RETRIES = 2
VIDEO_ID = re.compile(r'^[A-Za-z0-9_-]{11}$')


def canonical_url(raw):
    parsed = urlparse(str(raw).strip())
    if parsed.scheme != 'https' or parsed.username or parsed.password or parsed.port:
        raise ValueError('Chỉ chấp nhận URL HTTPS của video YouTube.')
    host = (parsed.hostname or '').lower()
    if host in ('youtu.be', 'www.youtu.be'):
        ident = parsed.path.strip('/')
    elif host in ('youtube.com', 'www.youtube.com', 'm.youtube.com'):
        path = parsed.path.strip('/').split('/')
        ident = (parse_qs(parsed.query).get('v') or [''])[0] if path == ['watch'] else (
            path[1] if len(path) == 2 and path[0] in ('shorts', 'live') else '')
    else:
        ident = ''
    if not VIDEO_ID.fullmatch(ident):
        raise ValueError('Link phải trỏ tới một video YouTube cụ thể (watch, shorts, live hoặc youtu.be).')
    return 'https://www.youtube.com/watch?v=' + ident


def pipeline_retryable_error(error):
    message = str(error).lower()
    return ('ai sửa thời lượng không hợp lệ' in message or
            'voice-repair-attempts' in message and 'access is denied' in message or
            'scene[' in message and any(detail in message for detail in
                                         ('không khớp mốc ảnh', 'end nhỏ hơn start',
                                          'đánh dấu point nhưng có end>start')))


def retryable_error(error):
    message = str(error).lower()
    if any(x in message for x in ('request too large', 'insufficient_quota', 'billing_hard_limit',
                                  'invalid api key', 'unauthorized', 'forbidden',
                                  'sign in to confirm you', 'xác minh phiên tải',
                                  'không đọc được cookie trình duyệt', 'thiếu cookies.txt',
                                  'winerror 10013', 'không được phép kết nối ra youtube')):
        return False
    if 'requested' in message and 'limit' in message and 'tokens per min' in message:
        return False
    if pipeline_retryable_error(message):
        return True
    return any(x in message for x in ('timed out', 'timeout', 'connection reset', 'connection aborted',
                                     'temporarily unavailable', 'too many requests', 'rate limit',
                                     '429', '502', '503', '504'))


def retry_policy(error, attempts):
    if not retryable_error(error):
        return 0, 0
    if pipeline_retryable_error(error):
        return MAX_AUTO_RETRIES, AUTO_RETRY_DELAY
    # Rate limits and service outages need breathing room, not rapid retries.
    return MAX_TRANSIENT_RETRIES, 15 * 3 ** attempts


def validate_items(items):
    if not 1 <= len(items) <= MAX_ITEMS:
        raise ValueError(f'Mỗi lô cần 1–{MAX_ITEMS} video.')
    result = []
    seen = set()
    for position, row in enumerate(items, 1):
        try:
            url = canonical_url(row.get('url', ''))
            if url in seen:
                raise ValueError('Link trùng trong danh sách.')
            seen.add(url)
            title = str(row.get('title') or '').strip()[:180]
            overrides = row.get('settings') or {}
            if not isinstance(overrides, dict) or set(overrides) - set(Settings.model_fields):
                raise ValueError('Cột cấu hình video không hợp lệ.')
            from .music import ASSET_FIELDS
            if set(overrides) & ASSET_FIELDS:
                raise ValueError('Chọn nhạc nền qua cấu hình chung, không nhập đường dẫn trong từng dòng.')
            if 'voice_reference' in overrides:
                raise ValueError('Chọn giọng mẫu qua cấu hình chung, không nhập đường dẫn trong từng dòng.')
            result.append({'url': url, 'title': title, 'settings': overrides, 'position': position})
        except (ValueError, AttributeError) as exc:
            raise ValueError(f'Dòng {position}: {exc}') from exc
    return result


def effective_settings(base, overrides=None):
    """Apply the same editorial-mode requirements as the single-video editor."""
    merged = {**base, **(overrides or {})}
    if merged.get('editorial_mode') == 'reaction_cops':
        merged.update(production_workflow='plan_first', narration_style='storytelling', opening_delay=0)
    return Settings.model_validate(merged).model_dump()


def parse_excel(content):
    """Read values only from the first sheet; never evaluate workbook formulas."""
    from openpyxl import load_workbook
    if len(content) > 2 * 1024 * 1024:
        raise ValueError('File Excel vượt giới hạn 2 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 15 * 1024 * 1024:
                raise ValueError('Dữ liệu bên trong file Excel vượt giới hạn 15 MB.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(x or '').strip().lower() for x in next(rows)]
        if 'url' not in headers:
            raise ValueError('Cột đầu vào bắt buộc tên url.')
        allowed = {'url', 'title', 'summary_minutes'} | set(Settings.model_fields) - {'voice_reference'}
        unknown = set(headers) - allowed - {''}
        if unknown or len(headers) != len(set(headers)):
            raise ValueError('Tên cột Excel không hợp lệ hoặc bị trùng: ' + ', '.join(sorted(unknown)))
        parsed = []
        for index, values in enumerate(rows, 2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if len(parsed) >= MAX_ITEMS:
                raise ValueError(f'File Excel chỉ được tối đa {MAX_ITEMS} video.')
            record = dict(zip(headers, values))
            overrides = {}
            for key, value in record.items():
                if key in ('', 'url', 'title') or value is None or str(value).strip() == '':
                    continue
                if key == 'summary_minutes':
                    overrides['summary_seconds'] = float(value) * 60
                elif isinstance(value, str):
                    try:
                        overrides[key] = json.loads(value)
                    except ValueError:
                        overrides[key] = value
                else:
                    overrides[key] = value
            parsed.append({'url': str(record.get('url') or '').strip(),
                           'title': str(record.get('title') or '').strip(),
                           'settings': overrides, 'row': index})
        workbook.close()
        return parsed
    except (OSError, IndexError, StopIteration, KeyError, zipfile.BadZipFile) as exc:
        raise ValueError('Không đọc được sheet đầu tiên của file Excel.') from exc


def save_reference(filename, content):
    ext = Path(filename or '').suffix.lower()
    if ext not in REFERENCE_EXTENSIONS:
        raise ValueError('Chọn file audio giọng mẫu.')
    if not content or len(content) > MAX_REFERENCE_BYTES:
        raise ValueError('Giọng mẫu phải có dữ liệu và tối đa 50 MB.')
    token = uuid.uuid4().hex
    folder = store.DATA / '_batch_references'
    folder.mkdir(parents=True, exist_ok=True)
    (folder / (token + ext)).write_bytes(content)
    return token + ext


def staged_reference(token):
    if not token or Path(token).name != token or len(Path(token).stem) != 32 or \
            any(c not in '0123456789abcdef' for c in Path(token).stem) or \
            Path(token).suffix.lower() not in REFERENCE_EXTENSIONS:
        raise ValueError('Giọng mẫu của lô không hợp lệ. Hãy tải lại file.')
    path = store.DATA / '_batch_references' / token
    if not path.is_file() or not 0 < path.stat().st_size <= MAX_REFERENCE_BYTES:
        raise ValueError('Giọng mẫu của lô không còn tồn tại. Hãy tải lại file.')
    return path


def create(name, items, settings, source_project_id=None, reference_token=None, music_token=None):
    rows = validate_items(items)
    base = Settings.model_validate(settings).model_dump()
    source = store.read(source_project_id) if source_project_id else None
    from . import music
    background_music = None
    if music_token:
        background_music, music_metadata = music.staged(music_token)
        base.update(music_metadata)
    elif base['music_file']:
        if not source or base['music_file'] != source['settings'].get('music_file'):
            raise ValueError('Nhạc nền cần file tải lên hoặc dự án nguồn tương ứng để sao chép an toàn.')
        background_music = music.project_asset(source)
        base.update({key: source['settings'].get(key, base[key]) for key in music.ASSET_FIELDS})
    if reference_token and base['voice_mode'] != 'clone':
        raise ValueError('Chỉ dùng giọng mẫu khi đã chọn Theo giọng mẫu.')
    if base['voice_reference'] and not source and not reference_token:
        raise ValueError('Giọng mẫu cần dự án nguồn để sao chép an toàn.')
    if source and base['voice_reference'] != source['settings'].get('voice_reference') and not reference_token:
        raise ValueError('Giọng mẫu không khớp dự án cấu hình đã chọn.')
    reference = staged_reference(reference_token) if reference_token else (store.asset(source_project_id, base['voice_reference']) if base['voice_reference'] else None)
    if reference and not reference.is_file():
        raise ValueError('File giọng mẫu của dự án nguồn không còn tồn tại.')
    if base['voice_mode'] == 'clone' and not reference:
        raise ValueError('Chọn file giọng mẫu tham chiếu trước khi tạo lô.')
    # Validate all overrides before creating any project.
    normalized = [effective_settings(base, row['settings']) for row in rows]
    from .story import output_budget
    for row, snapshot in zip(rows, normalized):
        output_budget(snapshot)
        if row['title']:
            snapshot['title'] = row['title']
        elif 'title' not in row['settings']:
            snapshot['title'] = ''
    batch_id = uuid.uuid4().hex
    created = time.time()
    with store.LOCK:
        projects = []
        try:
            for row, snapshot in zip(rows, normalized):
                project = store.create(row['title'] or 'Video YouTube mới',
                                       {'kind': 'youtube', 'url': row['url'], 'file': ''})
                projects.append(project)
                if reference:
                    target = store.asset(project['id'], reference.name)
                    if reference.resolve() != target.resolve():
                        shutil.copyfile(reference, target)
                    snapshot['voice_reference'] = target.name
                    if reference_token:
                        snapshot['voice_reference_text'] = ''
                        snapshot['voice_reference_hash'] = ''
                else:
                    snapshot['voice_reference'] = ''
                project['settings'] = snapshot
                if background_music:
                    target = store.asset(project['id'], background_music.name)
                    shutil.copyfile(background_music, target)
                    snapshot['music_file'] = target.name
                else:
                    snapshot.update(music_file='', music_name='', music_duration=0)
                project['batch_id'] = batch_id
                project['batch_position'] = row['position']
                store.save(project)  # Do not promote batch settings into user defaults.
            with store.conn() as db:
                db.execute('INSERT INTO batches VALUES (?,?,?,?,?)',
                           (batch_id, name.strip()[:180] or 'Lô video YouTube', 'running', created,
                            json.dumps(base, ensure_ascii=False)))
                db.executemany('INSERT INTO batch_items (id,batch_id,position,project_id,url,state,job_id,error,created,title) VALUES (?,?,?,?,?,?,?,?,?,?)',
                               [(uuid.uuid4().hex, batch_id, row['position'], project['id'], row['url'],
                                 'pending', None, '', created, row['title'] or 'Video YouTube mới') for row, project in zip(rows, projects)])
        except Exception:
            for project in projects:
                try:
                    store.delete_project(project['id'], project['revision'])
                except (OSError, RuntimeError, ValueError):
                    pass  # Preserve the original creation error.
            raise
    return get(batch_id)


def get(batch_id):
    with store.conn() as db:
        batch = db.execute('SELECT * FROM batches WHERE id=?', (batch_id,)).fetchone()
        if not batch:
            raise KeyError(batch_id)
        rows = db.execute('''SELECT bi.*,
                            j.progress AS progress, j.message AS message, j.error AS job_error,
                            ej.state AS export_state, ej.progress AS export_progress,
                            ej.error AS export_error, ej.message AS export_message
                            FROM batch_items bi
                            LEFT JOIN jobs j ON j.id=bi.job_id
                            LEFT JOIN jobs ej ON ej.id=bi.export_job_id
                            WHERE bi.batch_id=? ORDER BY bi.position''', (batch_id,)).fetchall()
    items = []
    for row in rows:
        items.append({'id': row['id'], 'position': row['position'], 'project_id': row['project_id'],
                      'url': row['url'], 'title': row['title'],
                      'state': row['state'],
                      'progress': 100 if row['state'] == 'completed' else row['progress'] or 0,
                      'message': row['message'] or '', 'error': row['error'] or row['job_error'] or '',
                      'exports': json.loads(row['exports']), 'job_id': row['job_id'],
                      'export_job_id': row['export_job_id'], 'export_folder': row['export_folder'],
                      'export_state': row['export_state'] or '',
                      'export_progress': row['export_progress'] or 0,
                      'export_error': row['export_error'] or '',
                      'export_message': row['export_message'] or '',
                      'attempts': row['attempts'], 'ready_at': row['ready_at'],
                      'retry_limit': retry_policy(row['error'] or row['job_error'], row['attempts'])[0]})
    return {'id': batch['id'], 'name': batch['name'], 'state': batch['state'],
            'created': batch['created'], 'items': items}


def list_batches():
    with store.conn() as db:
        ids = [row['id'] for row in db.execute('SELECT id FROM batches ORDER BY created DESC')]
    return [get(batch_id) for batch_id in ids]


def list_batch_overview():
    """Small list for the batch selector; details are loaded for one batch."""
    with store.conn() as db:
        return [dict(row) for row in db.execute(
            'SELECT id,name,state,created FROM batches ORDER BY created DESC')]


def control(batch_id, action, item_id=None):
    if action not in ('pause', 'resume', 'cancel', 'cancel_item', 'retry'):
        raise ValueError('Thao tác lô không hợp lệ.')
    with store.LOCK, store.conn() as db:
        batch = db.execute('SELECT * FROM batches WHERE id=?', (batch_id,)).fetchone()
        if not batch:
            raise KeyError(batch_id)
        if action == 'pause':
            db.execute("UPDATE batches SET state='paused' WHERE id=? AND state='running'", (batch_id,))
        elif action == 'resume':
            db.execute("UPDATE batches SET state='running' WHERE id=? AND state IN ('paused','completed')", (batch_id,))
        elif action == 'cancel':
            db.execute("UPDATE batches SET state='cancelled' WHERE id=?", (batch_id,))
            db.execute("UPDATE batch_items SET state='cancelled' WHERE batch_id=? AND state='pending'", (batch_id,))
            db.execute("UPDATE jobs SET cancelled=1 WHERE id IN (SELECT job_id FROM batch_items WHERE batch_id=? AND state IN ('queued','running'))", (batch_id,))
        elif action == 'cancel_item':
            row = db.execute('SELECT * FROM batch_items WHERE id=? AND batch_id=?', (item_id, batch_id)).fetchone()
            if not row or row['state'] not in ('pending', 'queued', 'running'):
                raise ValueError('Video này không còn đang chờ hoặc xử lý.')
            if row['state'] == 'pending':
                db.execute("UPDATE batch_items SET state='cancelled',error='Đã hủy trước khi bắt đầu.' WHERE id=?", (item_id,))
            else:
                db.execute('UPDATE jobs SET cancelled=1 WHERE id=?', (row['job_id'],))
        elif action == 'retry':
            row = db.execute('SELECT * FROM batch_items WHERE id=? AND batch_id=?', (item_id, batch_id)).fetchone()
            if not row or row['state'] not in ('failed', 'cancelled'):
                raise ValueError('Chỉ thử lại video đã lỗi hoặc bị hủy.')
            if not db.execute('SELECT 1 FROM projects WHERE id=?', (row['project_id'],)).fetchone():
                raise ValueError('Dự án nguồn đã bị xóa; hãy tạo lại từ URL.')
            db.execute("UPDATE batch_items SET state='pending',job_id=NULL,error='',attempts=0,ready_at=0 WHERE id=?", (item_id,))
            db.execute("UPDATE batches SET state='running' WHERE id=?", (batch_id,))
    return get(batch_id)


def tick(enqueue):
    """Reconcile completed jobs, then atomically claim pending batch items."""
    def schedule_retry(db, row, error):
        limit, wait = retry_policy(error, row['attempts'])
        if row['attempts'] >= limit:
            return
        retry_number = row['attempts'] + 1
        db.execute("UPDATE batch_items SET state='pending',job_id=NULL,error=?,attempts=?,ready_at=? WHERE id=?",
                   (f'Tự thử lại lượt {retry_number}/{limit} sau {wait}s. ' + error,
                    retry_number, time.time() + wait, row['id']))
        db.execute("UPDATE batches SET state='running' WHERE id=? AND state='completed_with_errors'",
                   (row['batch_id'],))

    with store.LOCK:
        with store.conn() as db:
            active = db.execute("""SELECT bi.*, j.state AS job_state, j.error AS job_error,
                                  b.state AS batch_state FROM batch_items bi
                                  JOIN batches b ON b.id=bi.batch_id
                                  LEFT JOIN jobs j ON j.id=bi.job_id
                                  WHERE bi.state IN ('queued','running')""").fetchall()
            for row in active:
                state = row['job_state']
                if state in ('queued', 'running'):
                    db.execute('UPDATE batch_items SET state=? WHERE id=?', (state, row['id']))
                elif state == 'completed':
                    from .media import probe
                    try:
                        project = store.read(row['project_id'])
                        exports = project.get('exports') or []
                        valid = bool(exports)
                        for export in exports:
                            path = store.asset(project['id'], export['file'])
                            info = probe(path)
                            valid = valid and path.suffix.lower() == '.mp4' and info['width'] > 0 and info['has_audio'] and info['duration'] > 0
                    except Exception:
                        exports = []
                        valid = False
                    if valid:
                        db.execute("UPDATE batch_items SET state='completed',error='',title=?,exports=? WHERE id=?",
                                   (project['name'], json.dumps(exports, ensure_ascii=False), row['id']))
                    else:
                        db.execute("UPDATE batch_items SET state='failed',error='File thành phẩm không đạt kiểm tra hình/tiếng hoặc không tồn tại.' WHERE id=?", (row['id'],))
                elif state == 'interrupted':
                    batch = db.execute('SELECT state FROM batches WHERE id=?', (row['batch_id'],)).fetchone()
                    if batch and batch['state'] != 'cancelled':
                        db.execute("UPDATE batch_items SET state='pending',job_id=NULL WHERE id=?", (row['id'],))
                    else:
                        db.execute("UPDATE batch_items SET state='cancelled' WHERE id=?", (row['id'],))
                elif (state == 'failed' and row['batch_state'] != 'cancelled' and
                      row['attempts'] < retry_policy(row['job_error'], row['attempts'])[0]):
                    schedule_retry(db, row, row['job_error'] or '')
                elif state in ('failed', 'cancelled'):
                    db.execute('UPDATE batch_items SET state=?,error=? WHERE id=?',
                               (state, row['job_error'] or '', row['id']))
                elif state is None:
                    db.execute("UPDATE batch_items SET state='failed',error='Không tìm thấy tác vụ đã lưu; hãy thử lại video này.' WHERE id=?", (row['id'],))
            # A batch created before this retry policy may already contain a
            # failed item. Adopt it once, including after an app restart.
            previous = db.execute("""SELECT bi.* FROM batch_items bi JOIN batches b ON b.id=bi.batch_id
                                     WHERE bi.state='failed' AND bi.attempts<?
                                       AND b.state IN ('running','completed_with_errors')""",
                                  (MAX_AUTO_RETRIES,)).fetchall()
            for row in previous:
                if row['attempts'] < retry_policy(row['error'], row['attempts'])[0]:
                    schedule_retry(db, row, row['error'])
            db.execute("UPDATE batches SET state='completed_with_errors' WHERE state='running' AND id NOT IN (SELECT batch_id FROM batch_items WHERE state IN ('pending','queued','running')) AND id IN (SELECT batch_id FROM batch_items WHERE state IN ('failed','cancelled'))")
            db.execute("UPDATE batches SET state='completed' WHERE state='running' AND id NOT IN (SELECT batch_id FROM batch_items WHERE state IN ('pending','queued','running'))")
            capacity = max(0, MAX_ACTIVE - db.execute("SELECT COUNT(*) FROM jobs WHERE state IN ('queued','running')").fetchone()[0])
            # Within a batch, the earliest unfinished video owns the lane,
            # including its retry cooldown. Other batches can still use the
            # remaining global worker capacity.
            pending = db.execute("""SELECT bi.* FROM batch_items bi
                JOIN batches b ON b.id=bi.batch_id
                WHERE b.state='running' AND bi.state='pending' AND bi.ready_at<=?
                  AND NOT EXISTS (SELECT 1 FROM batch_items earlier
                                  WHERE earlier.batch_id=bi.batch_id
                                    AND earlier.position<bi.position
                                    AND earlier.state IN ('pending','queued','running'))
                  AND NOT EXISTS (SELECT 1 FROM batch_items active
                                  WHERE active.batch_id=bi.batch_id
                                    AND active.state IN ('queued','running'))
                ORDER BY b.created,bi.position LIMIT ?""", (time.time(), capacity)).fetchall()
        for row in pending:
            with store.conn() as db:
                db.execute('BEGIN IMMEDIATE')
                batch = db.execute('SELECT state FROM batches WHERE id=?', (row['batch_id'],)).fetchone()
                item = db.execute('SELECT state,ready_at,position FROM batch_items WHERE id=?', (row['id'],)).fetchone()
                if (not batch or batch['state'] != 'running' or not item or item['state'] != 'pending'
                        or item['ready_at'] > time.time()):
                    continue
                if db.execute("""SELECT 1 FROM batch_items
                                 WHERE batch_id=? AND (
                                   (position<? AND state IN ('pending','queued','running'))
                                   OR state IN ('queued','running'))""",
                              (row['batch_id'], item['position'])).fetchone():
                    continue
                if db.execute("SELECT 1 FROM jobs WHERE project_id=? AND state IN ('queued','running')", (row['project_id'],)).fetchone():
                    continue
                jid = uuid.uuid4().hex
                db.execute('INSERT INTO jobs (id,project_id,kind,state,progress,message,error,created,cancelled,options) VALUES (?,?,?,?,?,?,?,?,0,?)',
                           (jid, row['project_id'], 'all', 'queued', 0, 'Đang chờ xử lý', '', time.time(), '{}'))
                db.execute("UPDATE batch_items SET state='queued',job_id=?,error='' WHERE id=?", (jid, row['id']))
            enqueue(jid)


def run_scheduler(stop, enqueue):
    while not stop.is_set():
        try:
            tick(enqueue)
        except Exception:
            import logging
            logging.getLogger(__name__).exception('Batch scheduler failed; will retry')
        stop.wait(1)
